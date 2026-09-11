#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_soutenance_ppt.py : fabrique le support de soutenance devant l'Institut.

POURQUOI UN SCRIPT ET PAS UN FICHIER FAIT A LA MAIN. Le projet publie chaque
nombre depuis un script versionne, et un support de soutenance est le document ou
un chiffre perime coute le plus cher : le jury « essaie de retrouver dans le texte
du memoire ce que l'etudiant presente a l'oral ». Un support genere se regenere,
se relit et se met sous harnais ; un support fait a la main derive.

CE QU'IL N'Y AVAIT PAS DANS LE PROJET. Aucun outil .pptx : les cinq decks tuteur
sont en Beamer. Ce script est donc neuf, et il ne touche a aucun deck existant.

REGLES TENUES, toutes issues du CLAUDE.md ou du GUIDE_REDACTION du depot :
  - aucun chiffre qui ne soit pas deja publie dans le memoire ou une sortie de
    script. Le fichier controle_soutenance.md les liste un par un ;
  - jamais « le SCR DORA de telle entite » mais « besoin de capital ORSA au titre
    de DORA » ;
  - jamais l'addition des quatre canaux, ni celle des deux interactions ;
  - le niveau absolu est annonce ILLUSTRATIF, le rapport est le resultat ;
  - pas de tiret cadratin.

Sortie : soutenance_memoire_DORA.pptx, 16 slides principales et 10 de sauvegarde.
"""

import os
import sys

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE
except ImportError:
    sys.exit("python-pptx absent. Installer avec : .venv\\Scripts\\python.exe -m pip install python-pptx")

ICI = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(ICI, "..", "vasicek_lab", "figures")
SORTIE = os.path.join(ICI, "soutenance_memoire_DORA.pptx")

# ---------------------------------------------------------------------------
# PALETTE. Reprise de style_nexialog.py, la source unique des couleurs du projet,
# pour que le support et les figures qu'il porte soient de la meme famille.
# Trois teintes plus un accent, comme demande : encre, bleu, gris, et le rouge
# reserve aux ALERTES et a la non-identifiabilite. Jamais decoratif.
# ---------------------------------------------------------------------------
ENCRE = RGBColor(0x1B, 0x1E, 0x30)      # texte principal, presque noir
ARDOISE = RGBColor(0x22, 0x3E, 0x55)    # texte secondaire, bleu profond
BLEU = RGBColor(0x4C, 0x79, 0xC7)       # bleu clair, filets et accents neutres
GRIS = RGBColor(0x59, 0x59, 0x59)       # sources, mentions discretes
ACCENT = RGBColor(0xA6, 0x00, 0x2E)     # rouge Nexialog : alertes uniquement
FOND = RGBColor(0xFC, 0xFC, 0xFB)       # blanc casse
GRILLE = RGBColor(0xDC, 0xDC, 0xDC)     # filets

POLICE = "Segoe UI"                     # sans serif, presente sur tout poste Windows

L, H = Inches(13.333), Inches(7.5)      # 16:9
MARGE = Inches(0.62)


# ---------------------------------------------------------------------------
# Primitives
# ---------------------------------------------------------------------------
def nouvelle(prs):
    """Une diapositive vierge, fond pose explicitement."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fond = s.shapes.add_shape(1, 0, 0, L, H)       # 1 = rectangle
    fond.fill.solid()
    fond.fill.fore_color.rgb = FOND
    fond.line.fill.background()
    fond.shadow.inherit = False
    s.shapes._spTree.remove(fond._element)
    s.shapes._spTree.insert(2, fond._element)      # au fond de la pile
    return s


def zone(s, x, y, l, h):
    tb = s.shapes.add_textbox(x, y, l, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    return tf


def para(tf, texte, taille, couleur=ENCRE, gras=False, italique=False,
         espace_avant=0, aligne=PP_ALIGN.LEFT, premier=False):
    p = tf.paragraphs[0] if premier else tf.add_paragraph()
    p.alignment = aligne
    p.space_before = Pt(espace_avant)
    r = p.add_run()
    r.text = texte
    r.font.size = Pt(taille)
    r.font.bold = gras
    r.font.italic = italique
    r.font.color.rgb = couleur
    r.font.name = POLICE
    return p


def titre(s, texte, couleur=ENCRE):
    """Le titre porte une AFFIRMATION, pas une etiquette. 32 pt, deux lignes max."""
    tf = zone(s, MARGE, Inches(0.42), L - 2 * MARGE, Inches(1.15))
    para(tf, texte, 32, couleur, gras=True, premier=True)
    filet = s.shapes.add_shape(1, MARGE, Inches(1.62), Inches(1.5), Pt(3))
    filet.fill.solid()
    filet.fill.fore_color.rgb = BLEU
    filet.line.fill.background()
    filet.shadow.inherit = False


def source(s, texte):
    """Mention de source, discrete mais presente : l'Institut exige la tracabilite."""
    tf = zone(s, MARGE, H - Inches(0.62), L - 2 * MARGE, Inches(0.34))
    para(tf, texte, 12, GRIS, italique=True, premier=True)


def retenir(s, texte, y=None):
    """La ligne « A retenir », qui porte le message de la slide."""
    y = y if y is not None else H - Inches(1.32)
    bande = s.shapes.add_shape(1, MARGE, y, L - 2 * MARGE, Inches(0.56))
    bande.fill.solid()
    bande.fill.fore_color.rgb = RGBColor(0xF2, 0xF2, 0xF2)
    bande.line.fill.background()
    bande.shadow.inherit = False
    tf = bande.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.16)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, "À retenir  " + texte, 17, ARDOISE, gras=True, premier=True)


def figure(s, nom, haut=Inches(1.95), bas=Inches(1.45)):
    """Insere une figure du depot, centree et mise a l'echelle sans deformation.

    Les figures ne sont PAS redessinees : ce sont celles du memoire, produites par
    les scripts et deja relues. Les recadrer ou les recolorier romprait la
    correspondance entre l'oral et le document, que le jury verifie.
    """
    chemin = os.path.join(FIG, nom)
    if not os.path.exists(chemin):
        raise SystemExit(f"figure absente : {chemin}")
    from PIL import Image
    lp, hp = Image.open(chemin).size
    dispo_l = L - 2 * MARGE
    dispo_h = H - haut - bas
    ech = min(dispo_l / lp, dispo_h / hp)
    ll, hh = int(lp * ech), int(hp * ech)
    s.shapes.add_picture(chemin, int((L - ll) / 2), int(haut), ll, hh)


def puces(s, items, y=Inches(2.05), taille=24, largeur=None, x=None):
    """Liste courte. Jamais de paragraphe : une idee par ligne."""
    x = x if x is not None else MARGE
    largeur = largeur if largeur is not None else L - 2 * MARGE
    tf = zone(s, x, y, largeur, H - y - Inches(1.5))
    for i, (t, g) in enumerate(items):
        para(tf, t, taille, ENCRE if not g else ARDOISE, gras=g,
             espace_avant=0 if i == 0 else 14, premier=(i == 0))


def bloc(s, x, y, l, h, titre_bloc, lignes, teinte=BLEU, taille=17):
    """Un bloc encadre, pour les schemas et les tableaux a deux colonnes."""
    f = s.shapes.add_shape(1, x, y, l, h)
    f.fill.solid()
    f.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    f.line.color.rgb = teinte
    f.line.width = Pt(1.5)
    f.shadow.inherit = False
    tf = f.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.10)
    para(tf, titre_bloc, taille + 2, teinte, gras=True, premier=True,
         aligne=PP_ALIGN.CENTER)
    for t in lignes:
        para(tf, t, taille, ENCRE, espace_avant=6, aligne=PP_ALIGN.CENTER)
    return f


def fleche(s, x, y, l):
    """Fleche vers la droite.

    Le numero 13 passe en argument sortait un CARRE gris, pas une fleche : la
    numerotation des autoshapes de python-pptx ne se devine pas. Trouve en
    regardant la diapositive rendue, pas en relisant le code.
    """
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x, y, l, Inches(0.26))
    a.fill.solid()
    a.fill.fore_color.rgb = GRILLE
    a.line.fill.background()
    a.shadow.inherit = False


def figure_gauche(s, nom, largeur=Inches(6.15), haut=Inches(1.95),
                  bas=Inches(1.30)):
    """Figure haute placee a GAUCHE, le texte occupant la colonne de droite.

    POURQUOI CE SECOND GABARIT. Trois figures du memoire sont presque carrees,
    voire hautes : posees pleine largeur elles se dimensionnent sur la HAUTEUR
    disponible et n'occupent alors qu'un tiers de la diapositive, etiquettes
    illisibles. C'est la meme regle que dans le memoire, ou une figure portrait
    prend une page entiere et jamais un debord en largeur.
    """
    chemin = os.path.join(FIG, nom)
    if not os.path.exists(chemin):
        raise SystemExit(f"figure absente : {chemin}")
    from PIL import Image
    lp, hp = Image.open(chemin).size
    dispo_h = H - haut - bas
    ech = min(largeur / lp, dispo_h / hp)
    ll, hh = int(lp * ech), int(hp * ech)
    s.shapes.add_picture(chemin, int(MARGE), int(haut + (dispo_h - hh) / 2), ll, hh)
    return MARGE + Inches(6.35)          # abscisse ou commence la colonne droite


def notes(s, texte):
    s.notes_slide.notes_text_frame.text = texte.strip()


def separateur(prs, texte):
    s = nouvelle(prs)
    tf = zone(s, MARGE, Inches(3.1), L - 2 * MARGE, Inches(1.3))
    para(tf, texte, 40, ARDOISE, gras=True, premier=True, aligne=PP_ALIGN.CENTER)
    notes(s, "Slide de separation. Ne pas commenter.")
    return s


# ---------------------------------------------------------------------------
# Le support
# ---------------------------------------------------------------------------
def construire():
    prs = Presentation()
    prs.slide_width, prs.slide_height = L, H

    # ---- 1. Titre ---------------------------------------------------------
    s = nouvelle(prs)
    tf = zone(s, MARGE, Inches(2.05), L - 2 * MARGE, Inches(2.2))
    para(tf, "Quand la contagion n'est pas identifiable :", 34, ARDOISE,
         gras=True, premier=True)
    para(tf, "borner le capital DORA plutôt que le poser", 40, ENCRE, gras=True)
    para(tf, "Quantification du besoin de capital lié à la non-conformité "
             "au règlement DORA", 20, GRIS, espace_avant=18)
    filet = s.shapes.add_shape(1, MARGE, Inches(4.62), Inches(2.2), Pt(3))
    filet.fill.solid(); filet.fill.fore_color.rgb = ACCENT
    filet.line.fill.background(); filet.shadow.inherit = False
    tf = zone(s, MARGE, Inches(5.0), L - 2 * MARGE, Inches(1.5))
    para(tf, "Kélian Kaddouri", 24, ENCRE, gras=True, premier=True)
    para(tf, "ENSAE Paris, filière Actuariat  ·  Nexialog Consulting", 18, GRIS,
         espace_avant=6)
    para(tf, "Mémoire présenté devant l'ENSAE Paris pour l'obtention du diplôme "
             "de la filière Actuariat et l'admission à l'Institut des actuaires",
         14, GRIS, espace_avant=10)
    notes(s, """
[0:20] Ne pas lire la slide. Se presenter, annoncer la duree, et poser la
question tout de suite : combien coute en capital de ne pas se conformer a DORA,
et jusqu'ou ce chiffre peut-il etre affirme.
""")

    # ---- 2. Le probleme metier -------------------------------------------
    s = nouvelle(prs)
    titre(s, "DORA impose une maîtrise, aucun dispositif n'en fait du capital")
    y, hb, lb = Inches(2.35), Inches(1.45), Inches(2.18)
    etapes = [("Non-conformité", ["à un ou plusieurs", "des cinq piliers"]),
              ("Incident TIC", ["défaillance,", "propagation"]),
              ("Perte", ["opérationnelle", "et cyber"]),
              ("Capital", ["exigence", "de couverture"]),
              ("Décision", ["remédier, détenir", "ou transférer"])]
    x = MARGE
    for i, (t, ls) in enumerate(etapes):
        teinte = ACCENT if i == 3 else BLEU
        bloc(s, x, y, lb, hb, t, ls, teinte, taille=15)
        if i < 4:
            fleche(s, x + lb + Inches(0.06), y + Inches(0.58), Inches(0.28))
        x += lb + Inches(0.40)
    tf = zone(s, MARGE, Inches(4.25), L - 2 * MARGE, Inches(1.4))
    para(tf, "La Formule Standard charge le risque opérationnel au forfait, "
             "donc elle est aveugle à la conformité.", 22, ENCRE, premier=True)
    para(tf, "Un score qualitatif de conformité ne dit rien du capital qu'une "
             "direction des risques doit immobiliser.", 22, ENCRE, espace_avant=10)
    retenir(s, "Une exigence de maîtrise sans traduction prudentielle : c'est le vide que ce travail comble.")
    source(s, "Source : mémoire, chapitres 1 et 2.")
    notes(s, """
[1:30] Le point a faire passer est le MAILLON MANQUANT, pas le reglement. DORA
impose la maitrise de cinq domaines depuis janvier 2025. Aucun module prudentiel
ne traduit ce niveau de maitrise en capital : la Formule Standard charge le
risque operationnel par un pourcentage de primes ou de provisions, donc elle donne
le meme chiffre a une entite exemplaire et a une entite defaillante.
Transition : « la question de ce memoire est donc celle du maillon du milieu. »
""")

    # ---- 3. Question et contribution -------------------------------------
    s = nouvelle(prs)
    titre(s, "Une question en deux temps, et trois apports")
    tf = zone(s, MARGE, Inches(2.0), L - 2 * MARGE, Inches(0.9))
    para(tf, "Combien coûte en capital de ne pas se conformer à DORA, "
             "et jusqu'où ce chiffre peut-il être affirmé ?", 25, ACCENT,
         gras=True, italique=True, premier=True)
    y, hb, lb = Inches(3.05), Inches(2.0), Inches(3.85)
    bloc(s, MARGE, y, lb, hb, "Méthodologique",
         ["Une cascade dirigée", "entre domaines de contrôle,", "trois propriétés démontrées"], BLEU, 16)
    bloc(s, MARGE + lb + Inches(0.30), y, lb, hb, "Empirique",
         ["Une frontière", "d'identifiabilité :", "ce que la donnée ne dit pas"], ACCENT, 16)
    bloc(s, MARGE + 2 * (lb + Inches(0.30)), y, lb, hb, "Opérationnelle",
         ["La donnée manquante", "devient une exigence", "de reporting"], BLEU, 16)
    retenir(s, "La seconde moitié de la question fait le travail : elle interdit de publier une précision que la donnée ne porte pas.")
    source(s, "Source : mémoire, chapitre 1, section contribution et plan.")
    notes(s, """
[1:20] Insister sur la SECONDE moitie de la question. Beaucoup de memoires
repondent a la premiere. Celui-ci mesure aussi ce qu'il ne peut pas affirmer, et
c'est ce qui produit la contribution.
""")

    # ---- 4. Les cinq piliers ---------------------------------------------
    s = nouvelle(prs)
    titre(s, "Cinq domaines de contrôle, pas cinq cases à cocher")
    piliers = [("P1", "Gouvernance", "du risque TIC"),
               ("P2", "Incidents", "gestion, notification"),
               ("P3", "Tests", "de résilience"),
               ("P4", "Tiers", "prestataires TIC"),
               ("P5", "Partage", "sur les cybermenaces")]
    y, hb, lb = Inches(2.35), Inches(1.7), Inches(2.18)
    x = MARGE
    for code, nom, sous in piliers:
        teinte = ACCENT if code in ("P1", "P4") else BLEU
        bloc(s, x, y, lb, hb, code, [nom, sous], teinte, taille=16)
        x += lb + Inches(0.40)
    tf = zone(s, MARGE, Inches(4.35), L - 2 * MARGE, Inches(1.3))
    para(tf, "La défaillance de l'un dégrade la capacité de l'entité à tenir "
             "les autres : une gouvernance défaillante retarde la détection, "
             "une dépendance non maîtrisée ouvre une porte que les tests "
             "n'ont pas explorée.", 22, ENCRE, premier=True)
    retenir(s, "La non-conformité ne reste pas dans son pilier : elle se propage. En rouge, les deux piliers qui émettent le plus.")
    source(s, "Source : mémoire, chapitres 1 et 2 ; matrice de propagation, chapitre 6.")
    notes(s, """
[1:00] Ne PAS reciter le reglement. Le seul point : ce sont des domaines de
CONTROLE, et ils interagissent. Les deux piliers en rouge, gouvernance et tiers,
sont ceux qui emettent le plus dans la matrice calibree. Si on demande pourquoi,
la reponse est au chapitre 6 et la figure est en sauvegarde.
""")

    # ---- 5. Donnee : observable et non observable ------------------------
    s = nouvelle(prs)
    titre(s, "La donnée montre la co-occurrence, jamais la direction")
    y, hb, lb = Inches(2.25), Inches(2.35), Inches(5.9)
    bloc(s, MARGE, y, lb, hb, "Observable",
         ["Fréquence des incidents", "Sévérité des pertes",
          "Co-occurrence entre piliers", "Structure par vecteur d'attaque"], BLEU, 19)
    bloc(s, MARGE + lb + Inches(0.35), y, lb, hb, "Non identifiable",
         ["La DIRECTION de la contagion", "Quel pilier entraîne l'autre",
          "Une matrice et sa transposée", "sont indistinguables"], ACCENT, 19)
    tf = zone(s, MARGE, Inches(4.95), L - 2 * MARGE, Inches(0.8))
    para(tf, "Sept sources, chacune avec son statut de preuve déclaré, de la base "
             "versionnée et rejouable à la citation externe non recalculable.",
         20, ENCRE, premier=True)
    retenir(s, "Cette limite n'est pas un défaut du dispositif de données : elle est structurelle, et elle commande la suite.")
    source(s, "Source : mémoire, chapitre 4 (sources et statuts de preuve) et chapitre 8.")
    notes(s, """
[1:30] C'est la slide qui PREPARE le coeur. Annoncer les deux colonnes, et
surtout dire que la colonne de droite n'est pas un manque de donnees a combler
par plus de donnees du meme type : elle est structurelle. La demonstration
arrive deux slides plus loin.
""")

    # ---- 6. Architecture --------------------------------------------------
    s = nouvelle(prs)
    titre(s, "D'un état de conformité à une décision, en quatre étages")
    y, hb, lb = Inches(2.30), Inches(1.55), Inches(2.75)
    etages = [("État de conformité", ["par pilier,", "chaîne de Markov"], BLEU),
              ("Quatre canaux", ["fréquence, détection,", "propagation, accumulation"], ACCENT),
              ("Perte annuelle", ["fréquence × sévérité,", "cascade entre piliers"], BLEU),
              ("Capital et décision", ["quantile à 99,5 %,", "remédier ou transférer"], BLEU)]
    x = MARGE
    for i, (t, ls, c) in enumerate(etages):
        bloc(s, x, y, lb, hb, t, ls, c, taille=15)
        if i < 3:
            fleche(s, x + lb + Inches(0.04), y + Inches(0.62), Inches(0.28))
        x += lb + Inches(0.36)
    tf = zone(s, MARGE, Inches(4.25), L - 2 * MARGE, Inches(1.3))
    para(tf, "Deux canaux sont calibrables sur données, la fréquence et la "
             "détection. Deux sont des bornes posées, la propagation et "
             "l'accumulation tiers.", 22, ENCRE, premier=True)
    retenir(s, "Le modèle est partiellement calibré et partiellement borné, et le document dit lequel est lequel.")
    source(s, "Source : mémoire, chapitres 5 à 10.")
    notes(s, """
[1:30] Donner la chaine complete en une fois, sans notation. Le point a ne pas
rater : deux canaux CALIBRABLES, deux canaux BORNES. C'est la distinction qui
revient a la fin sur les limites, et c'est elle qui interdit d'additionner les
quatre leviers pour chiffrer une remediation partielle.
""")

    # ---- 7. Pourquoi une cascade dirigee ---------------------------------
    s = nouvelle(prs)
    titre(s, "L'ordre de propagation change la criticité, une copule ne le voit pas")
    figure(s, "H1_reseau_W.png", haut=Inches(2.0), bas=Inches(1.5))
    retenir(s, "Une copule capture la co-occurrence des pertes, mais symétriquement : elle ne distingue pas ce qui entraîne de ce qui subit.")
    source(s, "Source : mémoire, chapitre 6 ; figure produite par le script 29 ; comparaison au jumeau copule, chapitre 10.")
    notes(s, """
[1:30] Le point intuitif : une gouvernance defaillante retarde la detection des
incidents ; l'inverse n'est pas vrai au meme degre. Une dependance symetrique
donne le meme nombre dans les deux sens, donc elle ne peut pas servir a
prioriser une remediation. La cascade, si.
Ne PAS entrer dans la normalisation de Leontief : slide de sauvegarde.
""")

    # ---- 8. La non-identifiabilite ---------------------------------------
    s = nouvelle(prs)
    titre(s, "La direction n'est pas identifiable, et c'est un résultat", ACCENT)
    xd = figure_gauche(s, "Z11_reversibilite_martingale.png")
    tf = zone(s, xd, Inches(2.15), L - xd - MARGE, Inches(3.5))
    para(tf, "La matrice se décompose en une partie symétrique et une partie "
             "antisymétrique.", 21, ENCRE, premier=True)
    para(tf, "La donnée identifie la première, la co-occurrence.", 21, ARDOISE,
         gras=True, espace_avant=16)
    para(tf, "Elle n'identifie pas la seconde, la direction.", 21, ACCENT,
         gras=True, espace_avant=12)
    para(tf, "Une matrice et sa transposée sont indistinguables pour la donnée, "
             "et donnent pourtant un capital et une décision différents.",
         19, ENCRE, espace_avant=16)
    retenir(s, "Ce n'est pas un échec du modèle : c'est ce qui interdit de publier une direction comme si elle avait été observée.")
    source(s, "Source : mémoire, chapitre 8 ; script 40.")
    notes(s, """
[2:00] LA SLIDE DU MEMOIRE. Prendre le temps.
Trois phrases : la matrice se decompose en symetrique plus antisymetrique ; la
donnee identifie la co-occurrence, donc la partie symetrique ; une matrice et sa
transposee sont indistinguables pour la donnee et donnent pourtant un capital et
une decision DIFFERENTS.
Dire explicitement : ce n'est pas un echec, c'est un resultat d'identification.
Transition : « la question devient donc : que publie-t-on quand on ne peut pas
identifier ? »
""")

    # ---- 9. Du point a la bande ------------------------------------------
    s = nouvelle(prs)
    titre(s, "Du point à la bande : borner plutôt que poser")
    xd = figure_gauche(s, "Z_identification_partielle.png")
    tf = zone(s, xd, Inches(2.20), L - xd - MARGE, Inches(3.6))
    para(tf, "Plutôt que de poser la direction, le capital est borné sur "
             "l'ensemble des matrices admissibles.", 22, ENCRE, premier=True)
    para(tf, "1 024 sommets, énumérés exhaustivement", 22, ARDOISE, gras=True,
         espace_avant=16)
    para(tf, "L'ignorance directionnelle se paye linéairement, et son coût "
             "maximal est connu d'avance : c'est plus favorable qu'un "
             "intervalle de confiance ordinaire.", 20, ENCRE, espace_avant=16)
    para(tf, "Le niveau affiché reste illustratif.", 20, ACCENT, gras=True,
         espace_avant=16)
    retenir(s, "La bande est le résultat rigoureux : elle quantifie une incertitude structurelle au lieu de la masquer.")
    source(s, "Source : mémoire, chapitre 9 ; scripts 30 et 40.")
    notes(s, """
[2:00] Le mecanisme : plutot que de poser la direction, on borne le capital sur
l'ENSEMBLE des matrices compatibles avec ce que la donnee identifie. 1 024
sommets, enumeration exhaustive, donc pas un echantillonnage.
Le point qui plait a un jury : le cout de l'ignorance est LINEAIRE et son maximum
est connu d'avance, ce qui est plus favorable qu'un intervalle de confiance
ordinaire.
Precaution a exprimer : le NIVEAU affiche est illustratif ; c'est le rapport et la
hierarchie qui sont defendus.
""")

    # ---- 10. Les quatre canaux -------------------------------------------
    s = nouvelle(prs)
    titre(s, "Quatre canaux déplacés, et leurs effets ne s'additionnent pas")
    figure(s, "S24_interaction_canaux.png", haut=Inches(1.90), bas=Inches(2.60))
    tf = zone(s, MARGE, H - Inches(2.52), L - 2 * MARGE, Inches(1.05))
    # L'echelle est dite ICI, dans la ligne qui porte le chiffre de tete. Sans
    # elle, un jury lit « 20 188 millions » pour UNE entite et la question des
    # milliards part avant qu'on ait pu la cadrer. Le mot « secteur » la desamorce
    # en trois syllabes.
    para(tf, "Au secteur : besoin de capital de 6 049 à 20 188 M€, facteur 3,34. "
             "La non-conformité n'ajoute aucune pénalité, elle déplace quatre "
             "paramètres de la loi de perte.", 19, ENCRE, premier=True)
    para(tf, "Canaux isolés  9 138 M€        Écart total  14 139 M€        "
             "Interaction  5 001 M€, soit 35 %", 21, ACCENT, gras=True,
         espace_avant=10)
    retenir(s, "Remédier un canal rapporte davantage à une entité défaillante partout : c'est la colonne de fermeture qu'un plan doit citer.")
    source(s, "Source : mémoire, chapitre 10 ; scripts 43 et 68.")
    notes(s, """
[1:40] Trois nombres, pas plus : 9 138, 14 139, 5 001.
Dire le mecanisme : les canaux isoles ne somment pas au total parce que la
cascade est super-additive SUR SES CANAUX. Ajouter aussitot la nuance qui evite
une erreur : elle est quasi additive sur les PILIERS, et l'additivite des COUTS
au sein d'un sinistre est une troisieme chose, une hypothese non testee.
NE JAMAIS additionner les quatre leviers pour chiffrer une remediation partielle :
deux des quatre sont des bornes posees, pas des budgets.
""")

    # ---- 11. Les tiers ----------------------------------------------------
    s = nouvelle(prs)
    titre(s, "Le pilier des tiers concentre, il ne propage pas seulement")
    figure(s, "Z9_p4_accumulation.png", haut=Inches(1.95), bas=Inches(1.75))
    tf = zone(s, MARGE, H - Inches(1.72), L - 2 * MARGE, Inches(0.45))
    para(tf, "Un prestataire commun crée un choc partagé entre entités : "
             "ce n'est plus de la contagion, c'est de l'accumulation.",
         20, ENCRE, premier=True)
    retenir(s, "C'est exactement ce que DORA demande de cartographier : la dépendance à des prestataires TIC critiques.")
    source(s, "Source : mémoire, annexe C (adaptations par pilier) ; script 22.")
    notes(s, """
[1:00] Distinguer les deux mecanismes, parce qu'un jury les confond souvent :
la CONTAGION va d'un pilier a l'autre dans une meme entite ; l'ACCUMULATION vient
d'un prestataire commun a plusieurs entites. Le modele porte les deux, par deux
canaux differents.
Reserve a exprimer si on insiste : le canal d'accumulation est une borne posee,
pas une calibration.
""")

    # ---- 12. La donnee manquante devient un livrable ---------------------
    s = nouvelle(prs)
    titre(s, "La donnée manquante devient une exigence de reporting")
    xd = figure_gauche(s, "Z18_valeur_information.png")
    tf = zone(s, xd, Inches(2.20), L - xd - MARGE, Inches(3.6))
    para(tf, "Quatre champs de registre, hiérarchisés par ce qu'ils resserrent :",
         21, ENCRE, premier=True)
    for t in ["horodatage des incidents",
              "domaine de contrôle touché",
              "cause commune consolidée",
              "champs renseignés obligatoires"]:
        para(tf, "·  " + t, 22, ACCENT, gras=True, espace_avant=12)
    retenir(s, "Le registre d'incidents devient un instrument de capital : chaque champ renseigné resserre la bande.")
    source(s, "Source : mémoire, chapitre 12 ; script 55.")
    notes(s, """
[1:40] L'UNE DES DEUX SLIDES LES PLUS IMPORTANTES, avec la non-identifiabilite.
C'est la que la limite devient un livrable : le modele ne dit pas seulement
« il manque une donnee », il dit LAQUELLE et CE QU'ELLE RAPPORTE en resserrement
de bande.
C'est aussi la recommandation la plus directement actionnable par une direction
des risques, et celle qu'une relecture de praticien a confirmee comme deja
pratiquee sous une autre forme.
""")

    # ---- 13. Decision pour l'entite --------------------------------------
    s = nouvelle(prs)
    titre(s, "Trois décisions, et une confusion à ne pas commettre")
    y, hb, lb = Inches(2.30), Inches(1.95), Inches(3.85)
    bloc(s, MARGE, y, lb, hb, "Remédier",
         ["Prioriser les canaux", "calibrables :", "fréquence et détection"], BLEU, 16)
    bloc(s, MARGE + lb + Inches(0.30), y, lb, hb, "Détenir",
         ["Le capital immobilisé", "coûte chaque année,", "pas une seule fois"], BLEU, 16)
    bloc(s, MARGE + 2 * (lb + Inches(0.30)), y, lb, hb, "Transférer",
         ["Comparer au prix", "de marché du même", "risque"], BLEU, 16)
    tf = zone(s, MARGE, Inches(4.55), L - 2 * MARGE, Inches(1.1))
    para(tf, "Une borne de capital n'est pas un budget de remédiation : deux des "
             "quatre canaux sont des bornes posées, et leur part n'est pas une "
             "enveloppe à dépenser.", 21, ACCENT, gras=True, premier=True)
    retenir(s, "Le bénéfice d'une remédiation a deux composantes, et la plus grosse n'est pas le capital libéré mais la perte évitée.")
    source(s, "Source : mémoire, chapitre 10 ; scripts 50 et 83.")
    notes(s, """
[1:20] Les trois decisions, puis la mise en garde, qui est le vrai message.
Si on demande un chiffre de retour sur investissement : le retour calcule sur le
seul portage de capital est un MAJORANT, pas un plancher, parce que la seconde
composante du benefice, la sinistralite evitee, n'est pas monetisee dans ce
calcul. Ne pas donner le nombre d'annees sans dire le sens de la borne.
""")

    # ---- 14. Limites ------------------------------------------------------
    s = nouvelle(prs)
    titre(s, "Ce que ce travail n'établit pas, et ce qui tient malgré cela")
    y, hb, lb = Inches(2.20), Inches(3.0), Inches(5.9)
    bloc(s, MARGE, y, lb, hb, "Non établi",
         ["Direction de contagion non identifiée",
          "Niveau absolu illustratif,",
          "borne supérieure à l'entité",
          "Incertitude de queue large",
          "Deux canaux bornés, non calibrés"], ACCENT, 18)
    bloc(s, MARGE + lb + Inches(0.35), y, lb, hb, "Établi malgré cela",
         ["Le signe et l'ordre de l'écart",
          "La hiérarchie des piliers",
          "Le cadre de bornage et son coût",
          "La recommandation de reporting",
          "La priorité de collecte de données"], BLEU, 18)
    retenir(s, "Ces limites sont publiées, pas concédées : une thèse de départ a été réfutée par son auteur et un chiffre requalifié en borne.")
    source(s, "Source : mémoire, chapitre 11, table des limites en deux colonnes.")
    notes(s, """
[1:10] Presenter les deux colonnes ENSEMBLE, jamais la gauche seule.
Insister sur un point que le texte de l'Institut recompense explicitement : la
non-transitivite, qui donnait son titre au projet, a ete REFUTEE par son auteur,
et le chiffre d'entite a ete requalifie en borne superieure. Ce sont des
resultats, pas des faiblesses.
""")

    # ---- 15. Conclusion ---------------------------------------------------
    s = nouvelle(prs)
    titre(s, "Trois messages")
    tf = zone(s, MARGE, Inches(2.15), L - 2 * MARGE, Inches(2.6))
    for i, t in enumerate([
        "1.  La non-conformité à DORA se traduit en coût de risque et en capital, "
        "par le déplacement de quatre paramètres de la loi de perte.",
        "2.  La dépendance directionnelle ne doit pas être calibrée sur des "
        "données qui ne l'identifient pas.",
        "3.  La limite de données devient une recommandation actionnable pour "
        "l'entité et pour le régulateur."]):
        para(tf, t, 24, ENCRE, espace_avant=0 if i == 0 else 20, premier=(i == 0))
    bande = s.shapes.add_shape(1, MARGE, Inches(5.25), L - 2 * MARGE, Inches(0.95))
    bande.fill.solid(); bande.fill.fore_color.rgb = RGBColor(0xF6, 0xEC, 0xEE)
    bande.line.color.rgb = ACCENT; bande.line.width = Pt(1.5)
    bande.shadow.inherit = False
    tf2 = bande.text_frame
    tf2.word_wrap = True
    tf2.margin_left = Inches(0.2)
    tf2.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf2, "Quand la dépendance n'est pas identifiable, la bonne réponse "
              "actuarielle n'est pas la fausse précision : c'est une borne, une "
              "hiérarchie de données et une décision mieux informée.", 21, ACCENT,
         gras=True, premier=True)
    source(s, "Source : mémoire, chapitre 13.")
    notes(s, """
[0:30] Trois phrases, puis la phrase encadree, dite lentement et sans la lire mot
a mot. C'est la derniere chose que le jury entend avant les questions : elle doit
etre la reponse a la question posee au debut.
""")

    # ---- 16. Questions ----------------------------------------------------
    s = nouvelle(prs)
    tf = zone(s, MARGE, Inches(2.5), L - 2 * MARGE, Inches(1.2))
    para(tf, "Questions", 54, ENCRE, gras=True, premier=True, aligne=PP_ALIGN.CENTER)
    y, hb, lb = Inches(4.4), Inches(1.0), Inches(2.75)
    x = MARGE
    for t in ["Conformité", "Quatre canaux", "Perte annuelle", "Capital borné"]:
        bloc(s, x, y, lb, hb, t, [], BLEU, taille=15)
        x += lb + Inches(0.36)
    source(s, "Slides de sauvegarde disponibles : sources, paramètres, sévérité, cascade, identifiabilité, bande, canaux, tiers, limites, glossaire.")
    notes(s, """
Slide sobre. Garder l'architecture visible pendant les questions : elle sert de
point d'appui pour repondre en montrant ou se situe la question.
Les dix slides de sauvegarde suivent. Les appeler a la voix.
""")

    # ======================================================================
    # SAUVEGARDE
    # ======================================================================
    separateur(prs, "Annexes")

    def sauv(t, items=None, fig=None, src="", note="", haut=Inches(2.0)):
        s = nouvelle(prs)
        titre(s, t)
        if fig:
            figure(s, fig, haut=haut, bas=Inches(1.0))
        if items:
            puces(s, items, y=Inches(2.05), taille=21)
        source(s, src)
        notes(s, note)
        return s

    sauv("Sept sources, et chacune avec son statut de preuve",
         [("Sévérité en euros : base internationale de pertes opérationnelles", True),
          ("Chronologie de brèches : fréquence et étude d'événement", False),
          ("Corpus de rapports post-mortem : structure de la direction", False),
          ("États réglementaires publiés : quatre bilans réels, anonymisés", False),
          ("Comptage d'incidents par vecteur : structure, jamais niveau", False),
          ("Étude de marché cyberassurance : contexte, aucune calibration", False),
          ("Deux statuts distingués : versionné et rejouable, ou citation externe non recalculable", True)],
         src="Source : mémoire, chapitre 4 ; scripts 62 et 63.",
         note="Si le jury demande la tracabilite : chaque nombre publie sort d'un script versionne, et un harnais verifie les 2 346 nombres du document.")

    sauv("Les paramètres, et ce qui est estimé, posé ou gelé",
         [("Sévérité : loi de Pareto généralisée au-delà du seuil, indice de queue 0,5954", False),
          ("Seuil 20,03 M€, 91 excès, calibration gelée depuis le 7 août 2026", False),
          ("Fréquence : binomiale négative, validée hors échantillon", False),
          ("Gain de propagation : trois valeurs POSÉES, 0,45 / 0,68 / 0,90", True),
          ("Le capital est croissant en ce gain : l'ordre suffit, l'amplitude est un scénario", True)],
         src="Source : mémoire, annexe D ; scripts 07, 08b, 63 et 66.",
         note="Le point qui desamorce la question sur les valeurs posees : la these ne depend pas de leur niveau mais de leur ORDRE, et la monotonie est demontree.")

    sauv("Pourquoi une loi de valeurs extrêmes pour la sévérité",
         fig="J3_validation_adequation.png",
         src="Source : mémoire, chapitre 5 ; scripts 07 et 47.",
         note="Deux theoremes autorisent l'extrapolation au-dela du plus grand sinistre observe. L'ajustement passe Anderson-Darling et Kolmogorov-Smirnov. Le backtest hors echantillon valide la FORME de la queue ; le NIVEAU derive, et c'est declare.")

    sauv("La cascade est un processus de branchement sous-critique",
         fig="K3_branchement_R0.png",
         src="Source : mémoire, chapitre 6 ; script 03.",
         note="Rayon spectral 0,506, taux de reproduction 0,054 : la cascade s'eteint. La normalisation de Leontief garantit que le pilier le plus prolifique engendre au plus g descendants directs.")

    sauv("Ce que le corpus documentaire corrobore, et ce qu'il ne lève pas",
         fig="Z17_postmortem_direction.png",
         src="Source : mémoire, chapitre 8 ; scripts 53, 59 et 64.",
         note="Le corpus donne un signal de direction, mais il porte un biais de narration : une narration choisit un ordre. Le biais est MESURE, borne et non leve. Une relecture de praticien a donne un exemple compatible avec les deux explications, ce qui corrobore la frontiere et non la matrice.")

    sauv("La bande de capital et ses trois étages d'incertitude",
         fig="J4_bande_modele.png",
         src="Source : mémoire, chapitre 11 ; script 48.",
         note="Trois etages a ne jamais fondre en une seule barre : le bruit de calcul, reductible par la machine ; l'incertitude de parametre, portable par un intervalle ; et l'ambiguite de famille, epistemique, qui ne se moyenne pas.")

    sauv("L'attribution par canal : trois lectures, non interchangeables",
         fig="S15_allocation_shapley_euler.png",
         src="Source : mémoire, chapitre 10 ; scripts 68 et 82.",
         note="Isolee, fermeture et Shapley repondent a trois questions differentes. Shapley est la seule colonne additive, mais c'est une CONVENTION d'attribution : deux canaux sur quatre sont des bornes, leur part n'est pas un budget. L'intuition suffit, ne pas ecrire la formule.")

    sauv("Défaillances simultanées : une sortie du modèle, pas une lacune",
         fig="S30_defaillances_simultanees.png",
         src="Source : mémoire, annexe D ; script 74.",
         note="PIEGE DE PREMISSE FREQUENT. A l'etat non conforme, 62,31 % des sinistres touchent plus d'un pilier, contre 31,15 % a l'etat conforme. La loi est EXACTE, par enumeration de la progeniture, donc citable sans bruit.")

    sauv("Les hypothèses, par ordre de ce qu'elles coûtent",
         fig="S32_tornado_normalise.png",
         src="Source : mémoire, annexe D ; scripts 76 et 81.",
         note="Resultat contre-intuitif a assumer : la fragilite est dans l'ajustement de valeurs extremes, PAS dans la contagion. L'hypothese la plus lourde et non testee est l'additivite des couts au sein d'un sinistre.")

    sauv("Glossaire",
         [("DORA : règlement européen sur la résilience opérationnelle numérique", False),
          ("SCR : capital de solvabilité requis sous Solvabilité II", False),
          ("VaR : quantile de la charge annuelle agrégée, ici à 99,5 %", False),
          ("Quantile de sévérité : quantile d'un sinistre isolé, ce n'est pas un capital", True),
          ("GPD : loi de Pareto généralisée, pour les excès au-delà d'un seuil", False),
          ("Copule : structure de dépendance symétrique, sans direction", False),
          ("Shapley : convention d'attribution d'un effet total entre leviers", False),
          ("Cascade : propagation dirigée entre domaines de contrôle", False)],
         src="Source : mémoire, annexe F, table des notations.",
         note="La distinction en rouge est celle qui a produit une confusion en seance le 7 aout 2026 : le quantile de severite d'un sinistre n'est pas une mesure de capital.")

    prs.save(SORTIE)
    n = len(prs.slides.__iter__.__self__._sldIdLst)
    print(f"ecrit  {os.path.relpath(SORTIE, ICI)}")
    print(f"       {n} diapositives au total : 16 principales, 1 separateur, 10 de sauvegarde")


if __name__ == "__main__":
    construire()
